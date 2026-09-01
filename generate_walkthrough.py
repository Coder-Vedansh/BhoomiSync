"""
Generate BhoomiSync_Prompt6_Walkthrough.docx
"""
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
import os

doc = Document()

# -- Styles --
style = doc.styles['Normal']
font = style.font
font.name = 'Calibri'
font.size = Pt(11)

# -- Title Page --
doc.add_paragraph()
doc.add_paragraph()
title = doc.add_heading('BhoomiSync — Prompt 6 Walkthrough', level=0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = subtitle.add_run('Authentication, Authorization, RBAC & Security Audit Management')
run.font.size = Pt(16)
run.font.color.rgb = RGBColor(16, 185, 129)

doc.add_paragraph()
meta = doc.add_paragraph()
meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
meta.add_run('BhoomiSync Cadastral Intelligence Platform\n').font.size = Pt(12)
meta.add_run('Implementation Date: August 2026\n').font.size = Pt(10)
meta.add_run('Document Version: 1.0').font.size = Pt(10)

doc.add_page_break()

# -- Table of Contents --
doc.add_heading('Table of Contents', level=1)
toc_items = [
    '1. Executive Summary',
    '2. Architecture Overview',
    '3. Database Models',
    '4. Authentication System',
    '5. Role-Based Access Control (RBAC)',
    '6. Security Audit Logging',
    '7. API Endpoints Reference',
    '8. Frontend Implementation',
    '9. Demo Accounts',
    '10. Automated Test Results',
    '11. Files Created & Modified',
    '12. Verification Evidence',
]
for item in toc_items:
    doc.add_paragraph(item, style='List Number')

doc.add_page_break()

# -- Section 1: Executive Summary --
doc.add_heading('1. Executive Summary', level=1)
doc.add_paragraph(
    'Prompt 6 implements a production-grade authentication, authorization, and security '
    'audit management layer for BhoomiSync. The system replaces the development-only '
    'X-User-Role header mechanism with server-side authenticated user identities, '
    'cryptographic JWT tokens, bcrypt password hashing, refresh token rotation with '
    'family-based reuse detection, progressive login lockout, and an immutable security '
    'audit log. The architecture supports four canonical roles (PUBLIC, SURVEYOR, '
    'GOVERNMENT_OFFICIAL, ADMIN) with granular permission-based authorization and is '
    'fully extensible for additional roles.'
)

doc.add_paragraph(
    'Key achievements:\n'
    '• 99/99 backend tests passing (24 new auth/RBAC/audit tests + 75 existing)\n'
    '• 0 TypeScript compilation errors in frontend build\n'
    '• Full backward compatibility with Prompts 1–5\n'
    '• 4 demo accounts seeded for instant testing\n'
    '• Comprehensive Security Admin dashboard with real-time session monitoring'
)

doc.add_page_break()

# -- Section 2: Architecture Overview --
doc.add_heading('2. Architecture Overview', level=1)
doc.add_paragraph(
    'The authentication architecture follows industry best practices for '
    'stateless API authentication with server-side session tracking:'
)

arch_table = doc.add_table(rows=8, cols=2)
arch_table.alignment = WD_TABLE_ALIGNMENT.CENTER
arch_table.style = 'Light Grid Accent 1'
headers = ['Component', 'Technology']
for i, h in enumerate(headers):
    arch_table.rows[0].cells[i].text = h

data = [
    ('Access Tokens', 'JWT (HS256) — 30-minute expiry'),
    ('Refresh Tokens', 'SHA-256 hashed, 7-day expiry, family rotation'),
    ('Password Hashing', 'Bcrypt with 12 work factor rounds'),
    ('Login Security', '5-attempt lockout, 15-minute cooldown'),
    ('Audit Logging', 'Immutable append-only with secret sanitization'),
    ('Frontend Auth', 'React Context + localStorage token management'),
    ('RBAC Engine', 'Permission-based, data-driven, no code changes for new roles'),
]
for i, (k, v) in enumerate(data):
    arch_table.rows[i + 1].cells[0].text = k
    arch_table.rows[i + 1].cells[1].text = v

doc.add_page_break()

# -- Section 3: Database Models --
doc.add_heading('3. Database Models', level=1)
doc.add_paragraph('All auth models are in backend/app/models/auth/models.py:')

models = [
    ('User', 'Core user identity with user_id, username, email, password_hash, is_active, is_verified, failed_login_attempts, locked_until.'),
    ('Role', 'Named role with description. Canonical: PUBLIC, SURVEYOR, GOVERNMENT_OFFICIAL, ADMIN.'),
    ('Permission', 'Granular permission string (e.g., parcel.read.public, land_record.import).'),
    ('user_roles', 'Many-to-many association table linking users to roles.'),
    ('role_permissions', 'Many-to-many association table linking roles to permissions.'),
    ('RefreshToken', 'Hashed refresh tokens with family_id for reuse detection, expiry, and revocation tracking.'),
    ('UserSession', 'Active login sessions with IP, user-agent, device info, and expiry.'),
    ('SecurityAuditLog', 'Immutable audit events with action, result, sanitized metadata, and timestamp.'),
]
for name, desc in models:
    p = doc.add_paragraph()
    run = p.add_run(f'{name}: ')
    run.bold = True
    p.add_run(desc)

doc.add_page_break()

# -- Section 4: Authentication System --
doc.add_heading('4. Authentication System', level=1)

doc.add_heading('4.1 Login Flow', level=2)
doc.add_paragraph(
    '1. Client sends POST /api/v1/auth/login with username_or_email and password.\n'
    '2. Server checks if account exists and is not locked.\n'
    '3. Bcrypt verifies password hash in constant time.\n'
    '4. On success: JWT access token (30min) + refresh token (7 days) issued.\n'
    '5. UserSession created with device fingerprint.\n'
    '6. LOGIN_SUCCESS audit event recorded.\n'
    '7. On failure: failed_login_attempts incremented. After 5 failures, account locked for 15 minutes.'
)

doc.add_heading('4.2 Token Refresh Flow', level=2)
doc.add_paragraph(
    '1. Client sends POST /api/v1/auth/refresh with refresh_token.\n'
    '2. Server computes SHA-256 hash and looks up token.\n'
    '3. If token is valid and not expired: new access + refresh tokens issued.\n'
    '4. Old refresh token is revoked.\n'
    '5. If revoked token is reused: ALL tokens in the family are invalidated, '
    'sessions terminated, TOKEN_REUSE_ATTACK_DETECTED logged.'
)

doc.add_heading('4.3 Password Security', level=2)
doc.add_paragraph(
    'Passwords are hashed with bcrypt (12 rounds). The password policy requires:\n'
    '• Minimum 8 characters\n'
    '• At least 1 uppercase, 1 lowercase, 1 digit, 1 special character\n'
    '• Passwords are NEVER logged, returned in responses, or stored in audit metadata'
)

doc.add_page_break()

# -- Section 5: RBAC --
doc.add_heading('5. Role-Based Access Control (RBAC)', level=1)

doc.add_paragraph('Permission matrix (subset):')
rbac_table = doc.add_table(rows=11, cols=5)
rbac_table.style = 'Light Grid Accent 1'
rbac_headers = ['Permission', 'PUBLIC', 'SURVEYOR', 'GOV_OFFICIAL', 'ADMIN']
for i, h in enumerate(rbac_headers):
    rbac_table.rows[0].cells[i].text = h

rbac_data = [
    ('survey.read', '✓', '✓', '✓', '✓'),
    ('dataset.upload', '✗', '✓', '✗', '✓'),
    ('parcel.read.public', '✓', '✓', '✓', '✓'),
    ('parcel.read.private', '✗', '✓', '✓', '✓'),
    ('parcel.verify', '✗', '✓', '✗', '✓'),
    ('land_record.import', '✗', '✗', '✓', '✓'),
    ('ai.infer', '✗', '✓', '✗', '✓'),
    ('user.read', '✗', '✗', '✓', '✓'),
    ('user.update', '✗', '✗', '✗', '✓'),
    ('audit.read', '✗', '✗', '✓', '✓'),
]
for i, row in enumerate(rbac_data):
    for j, val in enumerate(row):
        rbac_table.rows[i + 1].cells[j].text = val

doc.add_page_break()

# -- Section 6: Security Audit Logging --
doc.add_heading('6. Security Audit Logging', level=1)
doc.add_paragraph(
    'Every security-relevant action is recorded in an immutable, append-only audit log. '
    'All metadata is recursively sanitized — fields containing "password", "token", '
    '"secret", "authorization", or "credential" are replaced with ***REDACTED***.'
)

events = [
    'USER_REGISTERED', 'LOGIN_SUCCESS', 'LOGIN_FAILED', 'ACCOUNT_LOCKED',
    'LOGOUT', 'LOGOUT_ALL', 'TOKEN_ROTATED', 'TOKEN_REUSE_ATTACK_DETECTED',
    'PASSWORD_CHANGED', 'ROLE_ASSIGNED', 'ROLE_REMOVED', 'AUTHORIZATION_DENIED',
]
doc.add_paragraph('Recorded event types: ' + ', '.join(events))

doc.add_page_break()

# -- Section 7: API Endpoints --
doc.add_heading('7. API Endpoints Reference', level=1)

api_table = doc.add_table(rows=18, cols=4)
api_table.style = 'Light Grid Accent 1'
api_headers = ['Method', 'Endpoint', 'Description', 'Auth']
for i, h in enumerate(api_headers):
    api_table.rows[0].cells[i].text = h

endpoints = [
    ('POST', '/auth/register', 'Create new user', 'No'),
    ('POST', '/auth/login', 'Authenticate', 'No'),
    ('POST', '/auth/refresh', 'Rotate token', 'Token'),
    ('POST', '/auth/logout', 'End session', 'Yes'),
    ('POST', '/auth/logout-all', 'End all sessions', 'Yes'),
    ('GET', '/auth/me', 'Current profile', 'Yes'),
    ('POST', '/auth/change-password', 'Update password', 'Yes'),
    ('GET', '/auth/sessions', 'List sessions', 'Yes'),
    ('DELETE', '/auth/sessions/{id}', 'Revoke session', 'Yes'),
    ('GET', '/auth/users', 'List users', 'Admin'),
    ('GET', '/auth/users/{id}', 'User detail', 'Admin'),
    ('PUT', '/auth/users/{id}/status', 'Toggle status', 'Admin'),
    ('POST', '/auth/users/{id}/roles', 'Assign role', 'Admin'),
    ('DELETE', '/auth/users/{id}/roles/{r}', 'Remove role', 'Admin'),
    ('GET', '/auth/audit-logs', 'Query audit log', 'Admin'),
    ('GET', '/auth/security-stats', 'Security stats', 'Admin'),
    ('POST', '/auth/verify-account', 'Email verification', 'Token'),
]
for i, (method, endpoint, desc, auth) in enumerate(endpoints):
    api_table.rows[i + 1].cells[0].text = method
    api_table.rows[i + 1].cells[1].text = endpoint
    api_table.rows[i + 1].cells[2].text = desc
    api_table.rows[i + 1].cells[3].text = auth

doc.add_page_break()

# -- Section 8: Frontend Implementation --
doc.add_heading('8. Frontend Implementation', level=1)

fe_files = [
    ('auth/authTypes.ts', 'TypeScript interfaces for User, Role, Permission, Token, AuditLog, SecurityStats.'),
    ('auth/authApi.ts', 'Fetch-based API client with automatic token refresh rotation.'),
    ('auth/authStore.ts', 'Demo account presets and auth state types.'),
    ('auth/AuthContext.tsx', 'React Context provider for auth state, login, logout, permission checks.'),
    ('components/common/PermissionGuard.tsx', 'Component for gating UI features by permission/role.'),
    ('pages/LoginPage.tsx', 'Glassmorphism login form with 4 demo Quick-Fill cards and registration.'),
    ('pages/SecurityAdminPage.tsx', 'Security Admin dashboard: User Directory, Sessions, Audit Trail, RBAC Matrix.'),
    ('components/layout/Header.tsx', 'Updated with user profile pill and active role switcher dropdown.'),
    ('components/layout/Sidebar.tsx', 'Updated with Security Admin and Login navigation items.'),
    ('App.tsx', 'Wrapped in AuthProvider, added /login and /security-admin routes.'),
]
for path, desc in fe_files:
    p = doc.add_paragraph()
    run = p.add_run(f'{path}: ')
    run.bold = True
    p.add_run(desc)

doc.add_page_break()

# -- Section 9: Demo Accounts --
doc.add_heading('9. Demo Accounts', level=1)

demo_table = doc.add_table(rows=5, cols=4)
demo_table.style = 'Light Grid Accent 1'
demo_headers = ['Role', 'Email', 'Username', 'Password']
for i, h in enumerate(demo_headers):
    demo_table.rows[0].cells[i].text = h

demos = [
    ('ADMIN', 'admin@bhoomisync.demo', 'admin', 'AdminPassword@2026'),
    ('SURVEYOR', 'surveyor@bhoomisync.demo', 'surveyor', 'SurveyorPassword@2026'),
    ('GOVERNMENT_OFFICIAL', 'official@bhoomisync.demo', 'official', 'OfficialPassword@2026'),
    ('PUBLIC', 'citizen@bhoomisync.demo', 'citizen', 'CitizenPassword@2026'),
]
for i, (role, email, user, pw) in enumerate(demos):
    demo_table.rows[i + 1].cells[0].text = role
    demo_table.rows[i + 1].cells[1].text = email
    demo_table.rows[i + 1].cells[2].text = user
    demo_table.rows[i + 1].cells[3].text = pw

doc.add_page_break()

# -- Section 10: Automated Test Results --
doc.add_heading('10. Automated Test Results', level=1)
doc.add_paragraph(
    'Total backend test suite: 99/99 PASSED (0 failures, 0 errors)\n\n'
    'New Prompt 6 test files:\n'
    '• test_authentication.py — 13 tests (password hashing, JWT tokens, login, lockout, '
    'token rotation, reuse detection, profile, logout)\n'
    '• test_authorization.py — 7 tests (admin CRUD, citizen 403, surveyor verify, '
    'official import, privacy projection)\n'
    '• test_security_audit.py — 4 tests (secret sanitization, login audit events, '
    'audit log API, security stats API)\n\n'
    'Existing Prompt 1–5 tests: 75/75 PASSED (zero regression)\n\n'
    'Frontend: npm run build — 0 TypeScript compilation errors'
)

doc.add_page_break()

# -- Section 11: Files Created & Modified --
doc.add_heading('11. Files Created & Modified', level=1)

files_created = [
    'backend/app/models/auth/models.py',
    'backend/app/models/auth/__init__.py',
    'backend/app/modules/auth/password_service.py',
    'backend/app/modules/auth/token_service.py',
    'backend/app/modules/auth/permissions.py',
    'backend/app/modules/auth/audit_service.py',
    'backend/app/modules/auth/auth_schemas.py',
    'backend/app/modules/auth/dependencies.py',
    'backend/app/modules/auth/auth_service.py',
    'backend/app/modules/auth/auth_router.py',
    'backend/app/modules/auth/__init__.py',
    'backend/tests/test_authentication.py',
    'backend/tests/test_authorization.py',
    'backend/tests/test_security_audit.py',
    'frontend/src/auth/authTypes.ts',
    'frontend/src/auth/authApi.ts',
    'frontend/src/auth/authStore.ts',
    'frontend/src/auth/AuthContext.tsx',
    'frontend/src/components/common/PermissionGuard.tsx',
    'frontend/src/pages/LoginPage.tsx',
    'frontend/src/pages/SecurityAdminPage.tsx',
    'docs/AUTHENTICATION_ARCHITECTURE.md',
    'docs/RBAC_PERMISSION_MATRIX.md',
    'docs/SECURITY_GUIDE.md',
    'docs/AUDIT_LOGGING.md',
]
for f in files_created:
    doc.add_paragraph(f'[NEW] {f}', style='List Bullet')

files_modified = [
    'backend/app/models/__init__.py',
    'backend/app/core/config.py',
    'backend/app/api/router.py',
    'backend/app/db/seed.py',
    'backend/app/modules/land_records/routers/land_records_router.py',
    'backend/tests/conftest.py',
    'backend/tests/test_land_records.py',
    'frontend/src/App.tsx',
    'frontend/src/components/layout/Header.tsx',
    'frontend/src/components/layout/Sidebar.tsx',
    '.env.example',
]
for f in files_modified:
    doc.add_paragraph(f'[MODIFIED] {f}', style='List Bullet')

doc.add_page_break()

# -- Section 12: Verification Evidence --
doc.add_heading('12. Verification Evidence', level=1)
doc.add_paragraph(
    'Live API verification confirmed all endpoints operational:\n\n'
    '• POST /auth/login → 200 OK, JWT access token returned\n'
    '• GET /auth/security-stats → 200 OK, 4 total users, 4 active\n'
    '• GET /auth/users → 200 OK, 4 users (citizen, official, surveyor, admin)\n'
    '• GET /auth/audit-logs → 200 OK, LOGIN_SUCCESS and AUTHORIZATION_DENIED events\n'
    '• GET /auth/sessions → 200 OK, active JWT sessions listed\n'
    '• npm run build → 0 TypeScript compilation errors\n'
    '• pytest tests/ → 99/99 PASSED'
)

# -- Save --
output_path = r'd:\BhoomiSync\BhoomiSync_Prompt6_Walkthrough.docx'
doc.save(output_path)
print(f'Document saved to {output_path}')
