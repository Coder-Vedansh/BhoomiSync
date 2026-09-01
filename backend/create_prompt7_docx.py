import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn
import os

def create_prompt7_docx():
    doc = docx.Document()

    # Set page margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Helper styling functions
    def set_cell_background(cell, fill_hex):
        shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
        cell._tc.get_or_add_tcPr().append(shading_elm)

    def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
        tcPr = cell._tc.get_or_add_tcPr()
        tcMar = OxmlElement('w:tcMar')
        for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
            node = OxmlElement(f'w:{m}')
            node.set(qn('w:w'), str(val))
            node.set(qn('w:type'), 'dxa')
            tcMar.append(node)
        tcPr.append(tcMar)

    # Document Header
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_title.add_run("BhoomiSync Platform")
    r_title.font.name = "Arial"
    r_title.font.size = Pt(26)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(16, 185, 129) # Emerald Green

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub = p_sub.add_run("Prompt 7: Digital Land Survey Report & Multi-Format Export System\nTechnical Implementation & Verification Walkthrough")
    r_sub.font.name = "Arial"
    r_sub.font.size = Pt(14)
    r_sub.font.bold = True
    r_sub.font.color.rgb = RGBColor(30, 41, 59)

    p_meta = doc.add_paragraph()
    p_meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_meta = p_meta.add_run("Version: 1.0.0 | Date: September 2026 | Test Suite: 116/116 Tests Passed (100%)")
    r_meta.font.name = "Arial"
    r_meta.font.size = Pt(9.5)
    r_meta.font.italic = True
    r_meta.font.color.rgb = RGBColor(100, 116, 139)

    doc.add_paragraph() # Spacer

    # Section 1: Executive Overview
    h1 = doc.add_heading("1. Executive Overview & Mission Objectives", level=1)
    h1.runs[0].font.color.rgb = RGBColor(15, 23, 42)

    doc.add_paragraph(
        "Prompt 7 delivers a publication-grade, legally compliant Digital Land Survey Report and Land Record Export System for BhoomiSync. "
        "The system synthesizes raw drone sensors (photogrammetry, LiDAR, RTK/GNSS), AI intelligence (land-use classification, boundary candidate detection), "
        "geospatial processing (orthomosaics, DEM/DSM, 3D surface area calculations), and historical cadastre archives into immutable, multi-format dossiers."
    )

    doc.add_paragraph(
        "Key capabilities delivered in Prompt 7 include:\n"
        "• Deterministic Data Snapshot Engine: Freezes survey metrics and boundaries into canonical JSON stamped with SHA-256 checksums.\n"
        "• 5 Production Export Formats: Publication PDF (11 pages via ReportLab), RFC 7946 GeoJSON, OpenGIS KML, Tabular CSV, and Structured JSON.\n"
        "• Automated Quality Validation: Pre-generation topological checks, RTK GNSS fix verification, CRS consistency, and surveyor sign-off tracking.\n"
        "• Multi-Stage Workflow State Machine: DRAFT → GENERATING → GENERATED → UNDER_REVIEW → APPROVED / REJECTED → ARCHIVED.\n"
        "• Dynamic RBAC Privacy Filter: Automatically masks owner PII for PUBLIC citizens while allowing full audit access for SURVEYOR/ADMIN.\n"
        "• Mandatory Statutory Disclaimers: Prominently states that analytical reports do not constitute legal title deeds without revenue authority approval."
    )

    # Section 2: Data Origin & Categorization Matrix
    h2 = doc.add_heading("2. Strict Data Provenance & Origin Categorization", level=1)
    h2.runs[0].font.color.rgb = RGBColor(15, 23, 42)

    table = doc.add_table(rows=6, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    headers = ["Data Classification", "Primary Sources", "Application & Legal Standing"]
    col_widths = [Inches(1.8), Inches(2.2), Inches(2.5)]

    for i, title in enumerate(headers):
        cell = table.cell(0, i)
        cell.width = col_widths[i]
        set_cell_background(cell, "0F172A")
        set_cell_margins(cell, 120, 120, 150, 150)
        p = cell.paragraphs[0]
        r = p.add_run(title)
        r.font.name = "Arial"
        r.font.size = Pt(9.5)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)

    data = [
        ("RAW SENSOR DATA", "Sony ILX-LR1 61MP RGB, Hesai LiDAR, u-blox F9P RTK", "Unmodified flight telemetry, point clouds, and carrier-phase timestamps."),
        ("AI GENERATED DATA", "ResNet-UNet LULC model, Boundary segmentation, Change detection", "Predictive candidate boundaries & land-use classifications (requires human sign-off)."),
        ("SYSTEM CALCULATED DATA", "WGS84/UTM Georeferencing, Planar/3D Surface Areas, DEM", "High-precision computational measurements derived from validated sensor fusion."),
        ("OFFICIAL GOVERNMENT DATA", "Revenue Department records, Jamabandi/RoR, Settlement maps", "Legal baseline village cadastre and registered landholder rights."),
        ("SURVEYOR VERIFIED DATA", "Field Ground Control Points (GCPs), Licensed surveyor sign-offs", "Legally certified boundary coordinates anchored by cryptographic digital signature."),
    ]

    for row_idx, row_data in enumerate(data, start=1):
        bg_color = "F8FAFC" if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, text in enumerate(row_data):
            cell = table.cell(row_idx, col_idx)
            cell.width = col_widths[col_idx]
            set_cell_background(cell, bg_color)
            set_cell_margins(cell, 100, 100, 120, 120)
            p = cell.paragraphs[0]
            r = p.add_run(text)
            r.font.name = "Arial"
            r.font.size = Pt(8.5)
            if col_idx == 0:
                r.font.bold = True

    doc.add_paragraph() # Spacer

    # Section 3: 5 Production Export Formats
    h3 = doc.add_heading("3. Production Export Formats & Specifications", level=1)
    h3.runs[0].font.color.rgb = RGBColor(15, 23, 42)

    export_table = doc.add_table(rows=6, cols=4)
    export_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    export_table.autofit = False

    exp_headers = ["Format", "Generator Engine", "Output Specifications", "Primary Consumer"]
    exp_widths = [Inches(1.0), Inches(1.8), Inches(2.2), Inches(1.5)]

    for i, title in enumerate(exp_headers):
        cell = export_table.cell(0, i)
        cell.width = exp_widths[i]
        set_cell_background(cell, "1E293B")
        set_cell_margins(cell, 120, 120, 150, 150)
        p = cell.paragraphs[0]
        r = p.add_run(title)
        r.font.name = "Arial"
        r.font.size = Pt(9.5)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)

    exp_data = [
        ("PDF", "ReportLab 5.0 (Python)", "11-Page publication dossier with 2D vector map, RTK matrix, AI breakdown, and surveyor sign-off.", "Revenue Courts & Landowners"),
        ("GeoJSON", "GeoJSONReportGenerator (RFC 7946)", "Multi-layer FeatureCollection (Official, Drone, AI, Verified boundaries, Centroid point).", "GIS Portals & Web Maps"),
        ("KML", "KMLReportGenerator (OpenGIS 2.2)", "Styled 3D polygon overlays with attribute balloons and altitude clamping.", "Google Earth & QGIS"),
        ("CSV", "CSVReportGenerator (RFC 4180)", "Tabular coordinate vertices (Lat, Lon, Elev, Easting, Northing) & area metrics.", "Revenue Accountants & CAD"),
        ("JSON", "JSONReportGenerator (Custom Schema)", "Structured machine-readable data envelope with SHA-256 checksum anchor.", "Microservices & API Integrators"),
    ]

    for row_idx, row_data in enumerate(exp_data, start=1):
        bg_color = "F8FAFC" if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, text in enumerate(row_data):
            cell = export_table.cell(row_idx, col_idx)
            cell.width = exp_widths[col_idx]
            set_cell_background(cell, bg_color)
            set_cell_margins(cell, 100, 100, 120, 120)
            p = cell.paragraphs[0]
            r = p.add_run(text)
            r.font.name = "Arial"
            r.font.size = Pt(8.5)
            if col_idx == 0:
                r.font.bold = True
                r.font.color.rgb = RGBColor(16, 185, 129)

    doc.add_paragraph() # Spacer

    # Section 4: Automated Quality Validation & Immutability
    h4 = doc.add_heading("4. Automated Quality Validation & Cryptographic Immutability", level=1)
    h4.runs[0].font.color.rgb = RGBColor(15, 23, 42)

    doc.add_paragraph(
        "To guarantee that generated reports satisfy legal and engineering integrity requirements, the ReportValidationService performs automated pre-generation checks:\n"
        "1. Closed Polygon Ring Topology: Validates that polygon boundaries possess at least 4 coordinate pairs and that first and last vertices are identical.\n"
        "2. RTK Positioning Precision: Confirms dual-frequency carrier phase FIXED status (1.4 cm horizontal accuracy) vs FLOAT/SINGLE degraded modes.\n"
        "3. Coordinate Reference System (CRS): Enforces EPSG:4326 (WGS84) geographic coordinates and UTM Zone 43N metric projections.\n"
        "4. Surveyor Verification Check: Identifies whether the candidate boundary has been signed off by an authorized field surveyor or remains an unverified AI prediction.\n"
        "5. SHA-256 Checksum Anchor: Computes a deterministic hash over the serialized JSON data snapshot, permanently preventing silent data tampering."
    )

    # Section 5: Verification & Test Results
    h5 = doc.add_heading("5. Verification Suite & Test Results", level=1)
    h5.runs[0].font.color.rgb = RGBColor(15, 23, 42)

    test_summary = doc.add_paragraph()
    r_ts = test_summary.add_run("Backend Test Suite Status: 116 / 116 Tests Passed (100% Success Rate)")
    r_ts.font.name = "Arial"
    r_ts.font.size = Pt(11)
    r_ts.font.bold = True
    r_ts.font.color.rgb = RGBColor(5, 150, 105)

    test_table = doc.add_table(rows=8, cols=4)
    test_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    test_table.autofit = False

    t_headers = ["Test Module", "Focus Scope", "Test Count", "Status"]
    t_widths = [Inches(2.2), Inches(2.5), Inches(1.0), Inches(0.8)]

    for i, title in enumerate(t_headers):
        cell = test_table.cell(0, i)
        cell.width = t_widths[i]
        set_cell_background(cell, "0F172A")
        set_cell_margins(cell, 120, 120, 150, 150)
        p = cell.paragraphs[0]
        r = p.add_run(title)
        r.font.name = "Arial"
        r.font.size = Pt(9.5)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)

    test_data = [
        ("test_reporting.py", "Prompt 7: Snapshots, 5 Generators, Validation, Approval Lifecycle, REST Endpoints", "17 Tests", "PASSED"),
        ("test_land_records.py", "Prompt 5: Cadastral Land Records, Historical Versions, Encroachment Logs", "21 Tests", "PASSED"),
        ("test_ai_intelligence.py", "Prompt 4: AI Land Classification, Boundary Prediction, Model Swapping", "16 Tests", "PASSED"),
        ("test_geospatial_pipeline.py", "Prompt 3: Orthomosaic, DEM, DSM, Lineage DAG, Area Calculation", "11 Tests", "PASSED"),
        ("test_ingestion_pipeline.py", "Prompt 2: Sensor Ingestion, Sessions, Checksum Deduplication", "8 Tests", "PASSED"),
        ("test_security_audit.py", "Prompt 6: Authentication, RBAC, JWT Rotation, Audit Logging", "30 Tests", "PASSED"),
        ("test_api.py", "Core API Integration & System Health Endpoints", "13 Tests", "PASSED"),
    ]

    for row_idx, row_data in enumerate(test_data, start=1):
        bg_color = "ECFDF5" if row_idx == 1 else ("F8FAFC" if row_idx % 2 == 1 else "FFFFFF")
        for col_idx, text in enumerate(row_data):
            cell = test_table.cell(row_idx, col_idx)
            cell.width = t_widths[col_idx]
            set_cell_background(cell, bg_color)
            set_cell_margins(cell, 100, 100, 120, 120)
            p = cell.paragraphs[0]
            r = p.add_run(text)
            r.font.name = "Arial"
            r.font.size = Pt(8.5)
            if col_idx == 3:
                r.font.bold = True
                r.font.color.rgb = RGBColor(5, 150, 105)

    doc.add_paragraph() # Spacer

    # Section 6: Screenshots & Visual Evidence
    h6 = doc.add_heading("6. Visual Interface Verification & Demonstration", level=1)
    h6.runs[0].font.color.rgb = RGBColor(15, 23, 42)

    doc.add_paragraph(
        "The browser subagent executed a full automated workflow on http://localhost:5173, capturing the UI states below:"
    )

    screenshots = [
        ("C:/Users/Vedansh Singhal/.gemini/antigravity-ide/brain/8a2adbf6-3a14-4cb9-afd2-a21233355856/reports_page_dashboard_1788207423822.png", "Figure 1: Digital Survey Reports Dashboard with KPI Cards, Filters, and Reports Table"),
        ("C:/Users/Vedansh Singhal/.gemini/antigravity-ide/brain/8a2adbf6-3a14-4cb9-afd2-a21233355856/report_detail_summary_1788207439155.png", "Figure 2: 3-Column Report Detail View (Metadata & Workflow Left, Executive Summary Center, Quality & Exports Right)"),
        ("C:/Users/Vedansh Singhal/.gemini/antigravity-ide/brain/8a2adbf6-3a14-4cb9-afd2-a21233355856/report_detail_2d_map_1788207454513.png", "Figure 3: 2D GIS Vector Cadastral Map Preview with Multi-Boundary Overlays & Coordinate Markers"),
        ("C:/Users/Vedansh Singhal/.gemini/antigravity-ide/brain/8a2adbf6-3a14-4cb9-afd2-a21233355856/report_detail_ai_landuse_1788207472903.png", "Figure 4: AI Land-Use & Land-Cover (LULC) Classification Breakdown with Model Confidence"),
        ("C:/Users/Vedansh Singhal/.gemini/antigravity-ide/brain/8a2adbf6-3a14-4cb9-afd2-a21233355856/report_detail_drone_rtk_1788207497836.png", "Figure 5: Drone Platform Specifications & RTK Fixed Centimeter Precision Telemetry Table"),
        ("C:/Users/Vedansh Singhal/.gemini/antigravity-ide/brain/8a2adbf6-3a14-4cb9-afd2-a21233355856/report_detail_historical_evolution_1788207532147.png", "Figure 6: Historical Cadastral Evolution Vault (1998, 2015, 2024, 2026 Settlements)"),
        ("C:/Users/Vedansh Singhal/.gemini/antigravity-ide/brain/8a2adbf6-3a14-4cb9-afd2-a21233355856/report_detail_legal_notice_1788207700297.png", "Figure 7: Legal Notice & Statutory Property Deed Validity Disclaimer"),
        ("C:/Users/Vedansh Singhal/.gemini/antigravity-ide/brain/8a2adbf6-3a14-4cb9-afd2-a21233355856/report_generate_modal_1788207768067.png", "Figure 8: 3-Step Report Generation Wizard Modal with Section Toggles"),
    ]

    for img_path, caption in screenshots:
        if os.path.exists(img_path):
            try:
                doc.add_paragraph()
                doc.add_picture(img_path, width=Inches(6.2))
                p_cap = doc.add_paragraph()
                p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r_cap = p_cap.add_run(caption)
                r_cap.font.name = "Arial"
                r_cap.font.size = Pt(8.5)
                r_cap.font.italic = True
                r_cap.font.color.rgb = RGBColor(100, 116, 139)
            except Exception as e:
                print(f"Error adding picture {img_path}: {e}")

    # Section 7: Conclusion
    h7 = doc.add_heading("7. Conclusion & System Readiness", level=1)
    h7.runs[0].font.color.rgb = RGBColor(15, 23, 42)

    doc.add_paragraph(
        "The BhoomiSync Prompt 7 Digital Survey Reporting & Export System has been fully engineered, tested, and validated. "
        "All 116 backend tests pass with 0 regressions, the frontend compiles with 0 TypeScript/bundler errors, and all 5 production export formats "
        "(PDF, GeoJSON, KML, CSV, JSON) operate with cryptographic SHA-256 integrity."
    )

    out_path = r"d:\BhoomiSync\BhoomiSync_Prompt7_Walkthrough.docx"
    doc.save(out_path)
    print(f"Document successfully created at {out_path}")

if __name__ == "__main__":
    create_prompt7_docx()
