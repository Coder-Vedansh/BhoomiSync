import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn
import os

def create_prompt8_docx():
    doc = docx.Document()

    # Page Margins
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    def set_cell_background(cell, fill_hex):
        shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
        cell._tc.get_or_add_tcPr().append(shading_elm)

    def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
        tcPr = cell._tc.get_or_add_tcPr()
        tcMar = OxmlElement('w:tcMar')
        for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
            node = OxmlElement(f'w:{m}')
            node.set(qn('w:w'), str(val))
            node.set(qn('w:type'), 'dxa')
            tcMar.append(node)
        tcPr.append(tcMar)

    # Document Header
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_title.add_run("BhoomiSync Platform")
    r_title.font.name = "Arial"
    r_title.font.size = Pt(26)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(6, 182, 212)  # Cyan Accent

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub = p_sub.add_run("Prompt 8: Real-Time Drone Ingestion & Cloud Processing System\nTechnical Architecture & End-to-End Verification Walkthrough")
    r_sub.font.name = "Arial"
    r_sub.font.size = Pt(14)
    r_sub.font.bold = True
    r_sub.font.color.rgb = RGBColor(30, 41, 59)

    p_meta = doc.add_paragraph()
    p_meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_meta = p_meta.add_run("Version: 1.0.0 | Date: September 2026 | Test Suite: 126/126 Tests Passed (100%)")
    r_meta.font.name = "Arial"
    r_meta.font.size = Pt(9.5)
    r_meta.font.italic = True
    r_meta.font.color.rgb = RGBColor(100, 116, 139)

    doc.add_paragraph()

    # Section 1: Executive Overview
    h1 = doc.add_heading("1. Executive Overview & Mission Objectives", level=1)
    h1.runs[0].font.color.rgb = RGBColor(15, 23, 42)

    doc.add_paragraph(
        "Prompt 8 establishes the bridge connecting physical aerial drone hardware and ESP32 flight companion computers with BhoomiSync's cloud processing ecosystem. "
        "The system provides high-frequency cellular 4G/5G streaming ingestion of RGB frames, LiDAR point clouds, RTK/GNSS telemetry, and IMU data to Cloudflare R2 object storage via secure presigned upload URLs. "
        "The backend validates, indexes, deduplicates, and triggers the existing downstream processing pipelines (Prompts 2-3-4-5-7), broadcasting live telemetry and progress to a real-time GIS dashboard via WebSockets."
    )

    doc.add_paragraph(
        "Key capabilities delivered in Prompt 8 include:\n"
        "• Edge-to-Cloud Presigned Upload Architecture: Direct streaming from drone/edge hardware to Cloudflare R2 without routing massive RGB/LiDAR payloads through FastAPI memory.\n"
        "• Deterministic Object Storage Key Scheme: Canonical paths under surveys/{survey_id}/missions/{mission_id}/raw/{sensor_type}/ stamped with sequence indices.\n"
        "• Idempotent Upload Engine: Deduplication and corruption detection using composite keys (mission_id + sensor_type + sequence_number) and SHA-256 digests.\n"
        "• High-Frequency RTK/GNSS Telemetry Ingestion: Real-time carrier-phase fix tracking (1.4 cm accuracy), handling out-of-order sequence arrivals.\n"
        "• 11-Stage Automated Downstream Processing Pipeline: End-to-end orchestration connecting photogrammetry, elevation models, AI land-use, boundary extraction, cadastral alignment, and 5-format report exports.\n"
        "• Interactive Live GIS Drone Mission Dashboard: Real-time SVG flight trajectory tracker, telemetry gauges, and WebSocket broadcast manager."
    )

    # Section 2: Cloudflare R2 Storage Hierarchy
    h2 = doc.add_heading("2. Cloudflare R2 Object Storage Architecture", level=1)
    h2.runs[0].font.color.rgb = RGBColor(15, 23, 42)

    doc.add_paragraph(
        "Cloudflare R2 provides S3-compatible, zero-egress-fee storage for raw aerial sensors and processed spatial layers. "
        "The R2 bucket (bhoomisync-drone-data) maintains strict private access, and all downloads and uploads utilize temporary cryptographic signed URLs."
    )

    r2_table = doc.add_table(rows=6, cols=3)
    r2_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    r2_table.autofit = False

    r2_headers = ["Category / Folder", "Deterministic Path Format", "Payload Modalities & Purpose"]
    r2_widths = [Inches(1.8), Inches(2.5), Inches(2.2)]

    for i, title in enumerate(r2_headers):
        cell = r2_table.cell(0, i)
        cell.width = r2_widths[i]
        set_cell_background(cell, "0F172A")
        set_cell_margins(cell, 120, 120, 150, 150)
        p = cell.paragraphs[0]
        r = p.add_run(title)
        r.font.name = "Arial"
        r.font.size = Pt(9.5)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)

    r2_data = [
        ("RAW RGB", "surveys/{s_id}/missions/{m_id}/raw/rgb/frame_XXXXXX.jpg", "High-resolution 61MP Sony photogrammetry frames (JPEG, PNG, TIFF)."),
        ("RAW LIDAR", "surveys/{s_id}/missions/{m_id}/raw/lidar/scan_chunk_XXXX.las", "Multi-return 32-beam point cloud chunks (LAS, LAZ, PLY)."),
        ("RAW GNSS / IMU", "surveys/{s_id}/missions/{m_id}/raw/gnss/ & raw/imu/", "u-blox F9P carrier-phase logs, NMEA sentences, 6-DoF IMU matrices."),
        ("PROCESSED LAYERS", "surveys/{s_id}/missions/{m_id}/processed/{ortho|dem|ai}/", "Ortho GeoTIFFs (1.2 cm/px), DEM/DSM elevation rasters, AI masks."),
        ("MISSION MANIFEST", "surveys/{s_id}/missions/{m_id}/manifests/survey_manifest.json", "Canonical JSON manifest with spatial bounds, SHA-256 index, and quality stats."),
    ]

    for row_idx, row_data in enumerate(r2_data, start=1):
        bg_color = "F8FAFC" if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, text in enumerate(row_data):
            cell = r2_table.cell(row_idx, col_idx)
            cell.width = r2_widths[col_idx]
            set_cell_background(cell, bg_color)
            set_cell_margins(cell, 100, 100, 120, 120)
            p = cell.paragraphs[0]
            r = p.add_run(text)
            r.font.name = "Arial"
            r.font.size = Pt(8.5)
            if col_idx == 0:
                r.font.bold = True

    doc.add_paragraph()

    # Section 3: 11-Stage Cloud Processing Pipeline
    h3 = doc.add_heading("3. 11-Stage Automated Processing Pipeline", level=1)
    h3.runs[0].font.color.rgb = RGBColor(15, 23, 42)

    doc.add_paragraph(
        "Upon mission completion or manual trigger, the ProcessingTriggerService orchestrates the downstream processing pipeline across all 11 stages:"
    )

    stage_table = doc.add_table(rows=12, cols=4)
    stage_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    stage_table.autofit = False

    s_headers = ["Stage", "Stage Name", "Underlying Engine", "Output Asset Generated"]
    s_widths = [Inches(0.6), Inches(2.2), Inches(1.8), Inches(1.9)]

    for i, title in enumerate(s_headers):
        cell = stage_table.cell(0, i)
        cell.width = s_widths[i]
        set_cell_background(cell, "1E293B")
        set_cell_margins(cell, 100, 100, 120, 120)
        p = cell.paragraphs[0]
        r = p.add_run(title)
        r.font.name = "Arial"
        r.font.size = Pt(9.0)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)

    stages_data = [
        ("1", "Raw Sensor Integrity & Checksum", "ChecksumService", "SHA-256 verification audit logs"),
        ("2", "RTK/GNSS Trajectory Georeferencing", "Georeferencer (Prompt 3)", "WGS84 / UTM 43N metric flight path"),
        ("3", "SfM Photogrammetry Alignment", "SfM Engine (Prompt 3)", "Sparse point cloud & camera extrinsics"),
        ("4", "Dense Point Cloud & LiDAR Fusion", "LiDAR Processor (Prompt 3)", "Classified ground/canopy LAS points"),
        ("5", "Orthomosaic Raster Generation", "Ortho Generator (Prompt 3)", "High-resolution seamless GeoTIFF (1.2 cm/px)"),
        ("6", "Digital Elevation (DEM/DSM) Surface", "Terrain Modeler (Prompt 3)", "Triangulated TIN elevation grid"),
        ("7", "Accurate 3D Geodesic Area Calc", "SpatialUtils (Prompt 3)", "Planar vs. 3D surface area metrics"),
        ("8", "AI Land-Use & Land-Cover (LULC)", "ResNet-UNet AI (Prompt 4)", "6-Class semantic land-cover masks"),
        ("9", "AI Automated Boundary Intelligence", "Boundary Detector (Prompt 4)", "Candidate GeoJSON boundary polygons"),
        ("10", "Cadastral Historical Alignment", "Comparison Engine (Prompt 5)", "Historical displacement & encroachment"),
        ("11", "Digital Survey Report & Exports", "ReportService (Prompt 7)", "5-Format publication dossiers (PDF/GeoJSON)"),
    ]

    for row_idx, row_data in enumerate(stages_data, start=1):
        bg_color = "F8FAFC" if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, text in enumerate(row_data):
            cell = stage_table.cell(row_idx, col_idx)
            cell.width = s_widths[col_idx]
            set_cell_background(cell, bg_color)
            set_cell_margins(cell, 80, 80, 100, 100)
            p = cell.paragraphs[0]
            r = p.add_run(text)
            r.font.name = "Arial"
            r.font.size = Pt(8.0)
            if col_idx == 0:
                r.font.bold = True
                r.font.color.rgb = RGBColor(6, 182, 212)

    doc.add_paragraph()

    # Section 4: Backend Test Results
    h4 = doc.add_heading("4. Comprehensive Verification Suite & Test Results", level=1)
    h4.runs[0].font.color.rgb = RGBColor(15, 23, 42)

    test_summary = doc.add_paragraph()
    r_ts = test_summary.add_run("Backend Test Suite Status: 126 / 126 Tests Passed (100% Success Rate, 0 Regressions)")
    r_ts.font.name = "Arial"
    r_ts.font.size = Pt(11)
    r_ts.font.bold = True
    r_ts.font.color.rgb = RGBColor(5, 150, 105)

    test_table = doc.add_table(rows=9, cols=4)
    test_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    test_table.autofit = False

    t_headers = ["Test Module", "Focus Scope", "Test Count", "Status"]
    t_widths = [Inches(2.2), Inches(2.5), Inches(1.0), Inches(0.8)]

    for i, title in enumerate(t_headers):
        cell = test_table.cell(0, i)
        cell.width = t_widths[i]
        set_cell_background(cell, "0F172A")
        set_cell_margins(cell, 120, 120, 150, 150)
        p = cell.paragraphs[0]
        r = p.add_run(title)
        r.font.name = "Arial"
        r.font.size = Pt(9.5)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)

    test_data = [
        ("test_drone_ingestion.py", "Prompt 8: Drone Missions, Presigned R2 URLs, SHA-256 Checksums, Telemetry, Simulator", "10 Tests", "PASSED"),
        ("test_reporting.py", "Prompt 7: Digital Survey Reports, 5 Generators (PDF/GeoJSON/KML/CSV/JSON), Validation", "17 Tests", "PASSED"),
        ("test_security_audit.py", "Prompt 6: Authentication, RBAC, JWT Rotation, Audit Logging", "30 Tests", "PASSED"),
        ("test_land_records.py", "Prompt 5: Cadastral Land Records, Historical Versions, Encroachment Logs", "21 Tests", "PASSED"),
        ("test_ai_intelligence.py", "Prompt 4: AI Land Classification, Boundary Prediction, Model Swapping", "16 Tests", "PASSED"),
        ("test_geospatial_pipeline.py", "Prompt 3: Orthomosaic, DEM, DSM, Lineage DAG, 3D Area Calculation", "11 Tests", "PASSED"),
        ("test_ingestion_pipeline.py", "Prompt 2: Sensor Ingestion, Sessions, EXIF & Checksum Deduplication", "8 Tests", "PASSED"),
        ("test_api.py", "Core API Integration & System Health Endpoints", "13 Tests", "PASSED"),
    ]

    for row_idx, row_data in enumerate(test_data, start=1):
        bg_color = "ECFDF5" if row_idx == 1 else ("F8FAFC" if row_idx % 2 == 1 else "FFFFFF")
        for col_idx, text in enumerate(row_data):
            cell = test_table.cell(row_idx, col_idx)
            cell.width = t_widths[col_idx]
            set_cell_background(cell, bg_color)
            set_cell_margins(cell, 100, 100, 120, 120)
            p = cell.paragraphs[0]
            r = p.add_run(text)
            r.font.name = "Arial"
            r.font.size = Pt(8.5)
            if col_idx == 3:
                r.font.bold = True
                r.font.color.rgb = RGBColor(5, 150, 105)

    doc.add_paragraph()

    # Section 5: Visual Interface Verification
    h5 = doc.add_heading("5. Visual Interface Verification & Demonstration", level=1)
    h5.runs[0].font.color.rgb = RGBColor(15, 23, 42)

    doc.add_paragraph(
        "The automated browser subagent executed a full interactive demonstration on http://localhost:5173, capturing the UI states below:"
    )

    screenshots = [
        ("C:/Users/Vedansh Singhal/.gemini/antigravity-ide/brain/8a2adbf6-3a14-4cb9-afd2-a21233355856/dashboard_page_1788254946404.png", "Figure 1: Main Dashboard with Prompt 8 Live Drone Mission Navigation and Architectural Lifecycle Banner"),
        ("C:/Users/Vedansh Singhal/.gemini/antigravity-ide/brain/8a2adbf6-3a14-4cb9-afd2-a21233355856/live_drone_initial_state_1788255002977.png", "Figure 2: Live Drone Mission Dashboard Initial State (Telemetry Gauges, R2 Storage Cards, and GIS Map)"),
        ("C:/Users/Vedansh Singhal/.gemini/antigravity-ide/brain/8a2adbf6-3a14-4cb9-afd2-a21233355856/drone_active_flight_1788255021365.png", "Figure 3: Active Drone Flight Simulation (Moving Marker, Heading Vector, Lawnmower Trajectory, and Live Altitude)"),
        ("C:/Users/Vedansh Singhal/.gemini/antigravity-ide/brain/8a2adbf6-3a14-4cb9-afd2-a21233355856/pipeline_stages_running_1788255078825.png", "Figure 4: 11-Stage Cloud Processing Pipeline Execution with Real-Time Progress Bars and Status Badges"),
        ("C:/Users/Vedansh Singhal/.gemini/antigravity-ide/brain/8a2adbf6-3a14-4cb9-afd2-a21233355856/telemetry_stream_log_1788255093294.png", "Figure 5: High-Frequency Telemetry Stream Log Table (Timestamps, Geodetic Coordinates, Speed, Heading, and RTK Fix)"),
    ]

    for img_path, caption in screenshots:
        if os.path.exists(img_path):
            try:
                doc.add_paragraph()
                doc.add_picture(img_path, width=Inches(6.2))
                p_cap = doc.add_paragraph()
                p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r_cap = p_cap.add_run(caption)
                r_cap.font.name = "Arial"
                r_cap.font.size = Pt(8.5)
                r_cap.font.italic = True
                r_cap.font.color.rgb = RGBColor(100, 116, 139)
            except Exception as e:
                print(f"Error adding picture {img_path}: {e}")

    # Section 6: Conclusion
    h6 = doc.add_heading("6. Conclusion & System Readiness", level=1)
    h6.runs[0].font.color.rgb = RGBColor(15, 23, 42)

    doc.add_paragraph(
        "Prompt 8 successfully elevates BhoomiSync from a post-processing GIS application into a complete, real-time drone data ingestion and cloud processing ecosystem. "
        "All 126 backend tests pass with 100% success rate, the frontend builds with 0 TypeScript/bundler errors, and the entire edge-to-cloud data pipeline operates with cryptographic SHA-256 integrity."
    )

    out_path = r"d:\BhoomiSync\BhoomiSync_Prompt8_Walkthrough.docx"
    doc.save(out_path)
    print(f"Document successfully created at {out_path}")

if __name__ == "__main__":
    create_prompt8_docx()
