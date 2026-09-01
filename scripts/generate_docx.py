import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_background(cell, hex_color):
    """Sets cell background color."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Sets internal padding for table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def add_callout(doc, text, title="IMPORTANT NOTICE", border_color="059669", bg_color="F0FDF4"):
    """Adds a stylish callout box."""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, bg_color)
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:top w:val="none"/>
            <w:left w:val="single" w:sz="36" w:space="0" w:color="{border_color}"/>
            <w:bottom w:val="none"/>
            <w:right w:val="none"/>
        </w:tcBorders>
    ''')
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    run_title = p.add_run(f"📌 {title}\n")
    run_title.bold = True
    run_title.font.size = Pt(10.5)
    run_title.font.color.rgb = RGBColor(5, 150, 105)
    
    run_text = p.add_run(text)
    run_text.font.size = Pt(9.5)
    run_text.font.color.rgb = RGBColor(30, 41, 59)
    doc.add_paragraph()

def build_workflow_document(output_path):
    doc = Document()
    
    # Page setup - 1 inch margins
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        section.header.is_linked_to_previous = False
        
        # Header text
        hp = section.header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run("BhoomiSync | Cadastral Survey & Resurvey Platform Architecture")
        hrun.font.size = Pt(8.5)
        hrun.font.color.rgb = RGBColor(148, 163, 184)
        
        # Footer text
        fp = section.footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        frun = fp.add_run("BhoomiSync Technical Workflow Document • Confidential & Proprietary")
        frun.font.size = Pt(8.5)
        frun.font.color.rgb = RGBColor(148, 163, 184)

    # ---------------------------------------------------------------------------
    # COVER / TITLE SECTION
    # ---------------------------------------------------------------------------
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(20)
    title_p.paragraph_format.space_after = Pt(4)
    r_title = title_p.add_run("BHOOMISYNC")
    r_title.bold = True
    r_title.font.size = Pt(28)
    r_title.font.color.rgb = RGBColor(5, 150, 105)

    subtitle_p = doc.add_paragraph()
    subtitle_p.paragraph_format.space_after = Pt(14)
    r_sub = subtitle_p.add_run("Technology-Driven Rural Agricultural Land Survey, Resurvey, Mapping & Ingestion Platform")
    r_sub.bold = True
    r_sub.font.size = Pt(14)
    r_sub.font.color.rgb = RGBColor(30, 41, 59)

    meta_p = doc.add_paragraph()
    meta_p.paragraph_format.space_after = Pt(20)
    r_meta = meta_p.add_run("Drone Data Ingestion, Continuous Cloud Pipeline, Decoupled Gateways & GIS Workflow Specification\nVersion 2.0.0 • Foundation + Ingestion Pipeline Phase • Date: August 2026")
    r_meta.font.size = Pt(10)
    r_meta.font.color.rgb = RGBColor(100, 116, 139)

    # Callout: Strict Scope Restriction
    add_callout(
        doc,
        "DRONE FLIGHT CONTROL IS STRICTLY OUT OF SCOPE AND NOT IMPLEMENTED.\n"
        "BhoomiSync contains no MAVLink, ArduPilot, PX4, ROS 2, autonomous navigation, or waypoint controls. "
        "BhoomiSync is purely a Data Acquisition, Ingestion, Cloud Storage, AI Segmentation, and GIS Cadastral Platform.",
        title="MANDATORY ARCHITECTURAL SCOPE RESTRICTION",
        border_color="DC2626",
        bg_color="FEF2F2"
    )

    # ---------------------------------------------------------------------------
    # SECTION 1: EXECUTIVE SUMMARY & OBJECTIVE
    # ---------------------------------------------------------------------------
    h1 = doc.add_heading("1. Executive Summary & Project Objective", level=1)
    h1.paragraph_format.space_before = Pt(16)
    h1.paragraph_format.space_after = Pt(6)
    
    p = doc.add_paragraph(
        "BhoomiSync is engineered to resolve India's pressing rural land record challenges by modernizing agricultural "
        "cadastral surveying through drone sensor acquisition, continuous cloud ingestion, spatial photogrammetry, "
        "3D LiDAR terrain modeling, and AI-driven boundary segmentation. The system transforms high-resolution multi-modal "
        "sensor data into centimeter-accurate, georeferenced 2D land maps and verifiable agricultural parcel geometries."
    )
    p.paragraph_format.space_after = Pt(8)

    # ---------------------------------------------------------------------------
    # SECTION 2: HARDWARE-AGNOSTIC DATA INGESTION PIPELINE (PROMPT 2)
    # ---------------------------------------------------------------------------
    h2 = doc.add_heading("2. Drone Data Ingestion & Continuous Cloud Pipeline", level=1)
    h2.paragraph_format.space_before = Pt(16)
    h2.paragraph_format.space_after = Pt(6)

    doc.add_paragraph(
        "The Prompt 2 Ingestion Pipeline provides a robust, continuous streaming and batched ingestion architecture. "
        "The system accepts multi-sensor files, extracts rich spatial metadata, enforces cryptographic SHA-256 deduplication, "
        "maintains upload session lifecycles, and queues downstream processing jobs."
    )

    ingestion_tbl = doc.add_table(rows=4, cols=4)
    ingestion_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    i_headers = ["Sensor Payload", "Supported Formats", "Extracted Spatial Metadata", "Downstream Processing Job"]
    for i, h in enumerate(i_headers):
        cell = ingestion_tbl.cell(0, i)
        cell.paragraphs[0].text = h
        cell.paragraphs[0].runs[0].bold = True
        cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
        set_cell_background(cell, "0F172A")
        set_cell_margins(cell, 80, 80, 100, 100)

    i_rows = [
        ("Nadir Camera Imagery", "JPEG, PNG, TIFF, DNG", "GPS Lat/Lon/Alt, EXIF timestamp, Camera Make/Model, Focal Length, ISO", "IMAGE_PREPROCESSING"),
        ("3D LiDAR Point Cloud", "LAS, LAZ, PLY, PCD, CSV", "Point Count, 3D Bounding Box [X/Y/Z min & max], CRS, Elevation Span", "LIDAR_PREPROCESSING"),
        ("RTK / GNSS Positioning", "RINEX, NMEA ($GNGGA), JSON, CSV", "Fixed/Float RTK fix status, Horizontal/Vertical Precision (cm), Satellites", "GEOREFERENCING")
    ]

    for row_idx, row in enumerate(i_rows, start=1):
        for col_idx, text in enumerate(row):
            cell = ingestion_tbl.cell(row_idx, col_idx)
            cell.paragraphs[0].text = text
            cell.paragraphs[0].runs[0].font.size = Pt(8.5)
            set_cell_background(cell, "F8FAFC" if row_idx % 2 == 1 else "FFFFFF")
            set_cell_margins(cell, 70, 70, 90, 90)

    doc.add_paragraph()

    # ---------------------------------------------------------------------------
    # SECTION 3: UPLOAD SESSION & FILE LIFECYCLE STATE MACHINES
    # ---------------------------------------------------------------------------
    h3 = doc.add_heading("3. Upload Session & File Lifecycle State Machines", level=1)
    h3.paragraph_format.space_before = Pt(16)
    h3.paragraph_format.space_after = Pt(6)

    doc.add_paragraph(
        "To ensure high reliability over intermittent field cellular connections, the system models continuous uploads "
        "as state-machine driven sessions:\n"
        "• Session Lifecycles: CREATED ──> UPLOADING ──> COMPLETED / PARTIAL / FAILED / CANCELLED\n"
        "• File Lifecycles: PENDING ──> UPLOADING ──> UPLOADED ──> VALIDATED ──> QUEUED ──> PROCESSING ──> PROCESSED / FAILED\n"
        "• Deduplication Guarantee: Checks SHA-256 hashes against existing survey files to prevent redundant storage allocation."
    )

    # ---------------------------------------------------------------------------
    # SECTION 4: VISUAL PLATFORM WORKFLOW & SCREENSHOTS
    # ---------------------------------------------------------------------------
    h4 = doc.add_heading("4. Visual Platform Walkthrough & Verification", level=1)
    h4.paragraph_format.space_before = Pt(16)
    h4.paragraph_format.space_after = Pt(6)

    brain_dir = r"C:\Users\Vedansh Singhal\.gemini\antigravity-ide\brain\8a2adbf6-3a14-4cb9-afd2-a21233355856"

    screenshots = [
        ("dashboard_page_1788158221629.png", "Figure 1: BhoomiSync Operations Dashboard with Launch Drone Ingestion Station trigger."),
        ("gis_map_workbench_1788158666575.png", "Figure 2: Interactive GIS Map Workbench rendering real-time photo capture points, LiDAR footprint, and RTK flight path."),
        ("gis_map_popup_1788158902985.png", "Figure 3: Georeferenced Camera Photo popup inspector displaying altitude and coordinate HUD on satellite tiles."),
        ("provenance_dag_tab_1788155896665.png", "Figure 4: Immutable Provenance DAG Graph tracking Tiers 1-3 dataset transformations."),
        ("ai_inference_manifest_1788156129403.png", "Figure 5: AI Intelligence Workbench executing model-agnostic cadastral boundary segmentation inference."),
        ("historical_cadastre_page_1788156237208.png", "Figure 6: Historical Revenue Cadastre vs Drone Resurvey Comparison Workbench."),
        ("system_gateways_page_1788156356918.png", "Figure 7: System Architecture Monitor tracking ESP32PhoneGateway, MockStorageProvider, and PostGIS health.")
    ]

    for filename, caption in screenshots:
        full_img_path = os.path.join(brain_dir, filename)
        if os.path.exists(full_img_path):
            doc.add_paragraph()
            doc.add_picture(full_img_path, width=Inches(6.2))
            cp = doc.add_paragraph()
            cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            crun = cp.add_run(caption)
            crun.font.size = Pt(8.5)
            crun.font.italic = True
            crun.font.color.rgb = RGBColor(100, 116, 139)

    # ---------------------------------------------------------------------------
    # SECTION 5: REST API CATALOG (/api/v1/)
    # ---------------------------------------------------------------------------
    h5 = doc.add_heading("5. Canonical REST API Catalog (/api/v1/)", level=1)
    h5.paragraph_format.space_before = Pt(16)
    h5.paragraph_format.space_after = Pt(6)

    api_tbl = doc.add_table(rows=11, cols=3)
    api_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    a_headers = ["Method", "Canonical Endpoint (/api/v1/)", "Functional Description"]
    for i, h in enumerate(a_headers):
        cell = api_tbl.cell(0, i)
        cell.paragraphs[0].text = h
        cell.paragraphs[0].runs[0].bold = True
        cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
        set_cell_background(cell, "047857")
        set_cell_margins(cell, 80, 80, 100, 100)

    api_rows = [
        ("GET", "/api/v1/health", "Platform health, PostGIS status, storage provider, and active gateway type."),
        ("GET", "/api/v1/sensors", "List registered drone sensor payloads (Camera, LiDAR, RTK GNSS)."),
        ("POST", "/api/v1/surveys/{id}/upload-sessions", "Initialize continuous / batched ingestion sync session."),
        ("POST", "/api/v1/surveys/{id}/upload", "Multipart continuous sensor file upload with EXIF/LiDAR parsing & dedup."),
        ("GET", "/api/v1/files/{file_id}", "Retrieve stored file metadata, EXIF parameters, and storage keys."),
        ("POST", "/api/v1/files/{file_id}/retry", "Retry failed file validation or requeue downstream processing job."),
        ("GET", "/api/v1/surveys/{id}/spatial-footprint", "GeoJSON FeatureCollections of photo points, LiDAR bounds, RTK tracks."),
        ("GET", "/api/v1/processing/jobs", "Monitor modular processing job queue."),
        ("GET", "/api/v1/surveys", "List all agricultural survey missions with pagination and summary stats."),
        ("GET", "/api/v1/surveys/{id}/parcels", "Retrieve all parcel boundaries with calculated geodesic area & perimeter.")
    ]

    for row_idx, row in enumerate(api_rows, start=1):
        for col_idx, text in enumerate(row):
            cell = api_tbl.cell(row_idx, col_idx)
            cell.paragraphs[0].text = text
            cell.paragraphs[0].runs[0].font.size = Pt(8.5)
            if col_idx == 0:
                cell.paragraphs[0].runs[0].bold = True
            set_cell_background(cell, "F0FDF4" if row_idx % 2 == 1 else "FFFFFF")
            set_cell_margins(cell, 60, 60, 80, 80)

    doc.add_paragraph()

    # ---------------------------------------------------------------------------
    # SECTION 6: TEST & VERIFICATION REPORT
    # ---------------------------------------------------------------------------
    h6 = doc.add_heading("6. Verification & Automated Test Results", level=1)
    h6.paragraph_format.space_before = Pt(16)
    h6.paragraph_format.space_after = Pt(6)

    doc.add_paragraph(
        "• Pytest Test Suite: 27/27 unit and integration tests passed in 0.95 seconds.\n"
        "• Frontend Production Build: Vite & TypeScript build completed in 1.45s with 0 compilation errors.\n"
        "• Browser End-to-End Testing: Verified Drone Ingestion Station, Demo Simulation Mode batch injection, and GIS Map spatial footprint layers."
    )

    doc.save(output_path)
    print(f"Document successfully generated at: {output_path}")

if __name__ == "__main__":
    out1 = r"d:\BhoomiSync\BhoomiSync_Workflow_Document_v2.docx"
    out2 = r"d:\BhoomiSync\docs\BhoomiSync_Workflow_Document.docx"
    os.makedirs(os.path.dirname(out2), exist_ok=True)
    try:
        build_workflow_document(r"d:\BhoomiSync\BhoomiSync_Workflow_Document.docx")
    except PermissionError:
        print("Note: Primary file is open in MS Word, generating BhoomiSync_Workflow_Document_v2.docx")
        build_workflow_document(out1)
    try:
        build_workflow_document(out2)
    except PermissionError:
        pass
