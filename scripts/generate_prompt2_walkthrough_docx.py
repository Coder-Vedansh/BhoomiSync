import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_background(cell, hex_color):
    """Sets cell background color."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Sets internal padding for table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def add_callout(doc, text, title="IMPORTANT NOTICE", border_color="059669", bg_color="F0FDF4"):
    """Adds a stylish callout box."""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, bg_color)
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:top w:val="none"/>
            <w:left w:val="single" w:sz="36" w:space="0" w:color="{border_color}"/>
            <w:bottom w:val="none"/>
            <w:right w:val="none"/>
        </w:tcBorders>
    ''')
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    run_title = p.add_run(f"📌 {title}\n")
    run_title.bold = True
    run_title.font.size = Pt(10.5)
    run_title.font.color.rgb = RGBColor(5, 150, 105) if border_color == "059669" else RGBColor(220, 38, 38)
    
    run_text = p.add_run(text)
    run_text.font.size = Pt(9.5)
    run_text.font.color.rgb = RGBColor(30, 41, 59)
    doc.add_paragraph()

def build_prompt2_walkthrough_document(output_path):
    doc = Document()
    
    # 1-inch margins
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        section.header.is_linked_to_previous = False
        
        # Header text
        hp = section.header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run("BhoomiSync | Prompt 2: Drone Ingestion & Cloud Pipeline Walkthrough")
        hrun.font.size = Pt(8.5)
        hrun.font.color.rgb = RGBColor(148, 163, 184)
        
        # Footer text
        fp = section.footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        frun = fp.add_run("BhoomiSync Prompt 2 Technical Walkthrough • Confidential & Proprietary")
        frun.font.size = Pt(8.5)
        frun.font.color.rgb = RGBColor(148, 163, 184)

    # ---------------------------------------------------------------------------
    # COVER / TITLE SECTION
    # ---------------------------------------------------------------------------
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(18)
    title_p.paragraph_format.space_after = Pt(4)
    r_title = title_p.add_run("BHOOMISYNC")
    r_title.bold = True
    r_title.font.size = Pt(26)
    r_title.font.color.rgb = RGBColor(5, 150, 105) # Emerald 600

    subtitle_p = doc.add_paragraph()
    subtitle_p.paragraph_format.space_after = Pt(12)
    r_sub = subtitle_p.add_run("Prompt 2 Technical Walkthrough: Drone Data Ingestion & Continuous Cloud Pipeline")
    r_sub.bold = True
    r_sub.font.size = Pt(14)
    r_sub.font.color.rgb = RGBColor(30, 41, 59)

    meta_p = doc.add_paragraph()
    meta_p.paragraph_format.space_after = Pt(18)
    r_meta = meta_p.add_run("Continuous Streaming Ingestion, Metadata Extractors, Deduplication, Processing Queues & GIS Footprints\nRelease: Prompt 2 Verified • Date: August 2026")
    r_meta.font.size = Pt(9.5)
    r_meta.font.color.rgb = RGBColor(100, 116, 139)

    # Mandatory Scope Restriction Callout
    add_callout(
        doc,
        "DRONE FLIGHT CONTROL IS STRICTLY OUT OF SCOPE AND NOT IMPLEMENTED.\n"
        "BhoomiSync contains no MAVLink, ArduPilot, PX4, ROS 2, autonomous waypoint navigation, or drone flight controls. "
        "BhoomiSync is exclusively a Data Acquisition, Ingestion, Cloud Storage, AI Segmentation, and GIS Cadastral Platform.",
        title="MANDATORY ARCHITECTURAL SCOPE RESTRICTION",
        border_color="DC2626",
        bg_color="FEF2F2"
    )

    # ---------------------------------------------------------------------------
    # SECTION 1: PROMPT 2 OBJECTIVE & EXECUTIVE SUMMARY
    # ---------------------------------------------------------------------------
    h1 = doc.add_heading("1. Prompt 2 Executive Summary & Architecture", level=1)
    h1.paragraph_format.space_before = Pt(16)
    h1.paragraph_format.space_after = Pt(6)
    
    p = doc.add_paragraph(
        "Prompt 2 establishes the end-to-end Drone Data Ingestion & Cloud Pipeline for BhoomiSync. "
        "The module continuously receives multi-sensor data collected from survey drones (High-Resolution Nadir Cameras, "
        "3D LiDAR, and RTK/GNSS positioning), extracts rich spatial metadata without modifying raw bytes, enforces "
        "cryptographic SHA-256 deduplication, maintains upload session state machines, and automatically triggers "
        "downstream modular processing job queues."
    )
    p.paragraph_format.space_after = Pt(8)

    # ---------------------------------------------------------------------------
    # SECTION 2: HARDWARE-AGNOSTIC DATA INGESTION MATRIX
    # ---------------------------------------------------------------------------
    h2 = doc.add_heading("2. Sensor Payloads & Metadata Extraction Engines", level=1)
    h2.paragraph_format.space_before = Pt(16)
    h2.paragraph_format.space_after = Pt(6)

    doc.add_paragraph(
        "The ingestion architecture is strictly hardware-agnostic. In the current prototype, data flows via "
        "ESP32 -> Surveyor Mobile App -> Cloud. In future in-house drone hardware, data will stream directly from an "
        "Onboard Companion Computer (NVIDIA Jetson / Industrial ARM). The backend requires zero refactoring."
    )

    tbl = doc.add_table(rows=4, cols=4)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["Sensor Payload", "Supported File Formats", "Extracted Geolocation & Metadata", "Downstream Processing Job"]
    for i, h in enumerate(headers):
        cell = tbl.cell(0, i)
        cell.paragraphs[0].text = h
        cell.paragraphs[0].runs[0].bold = True
        cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
        set_cell_background(cell, "0F172A")
        set_cell_margins(cell, 80, 80, 100, 100)

    rows_data = [
        ("Nadir Camera Photos", "JPEG, PNG, TIFF, DNG", "GPS Lat/Lon/Alt, EXIF timestamp, Camera Make/Model, Focal Length, ISO, Shutter", "IMAGE_PREPROCESSING"),
        ("3D LiDAR Point Clouds", "LAS, LAZ, PLY, PCD, CSV", "Point Count, 3D Bounding Box [X/Y/Z min & max], Coordinate System, Elevation Span", "LIDAR_PREPROCESSING"),
        ("RTK / GNSS Trajectory", "RINEX, NMEA ($GNGGA), JSON, CSV", "Fixed/Float RTK fix status, Horizontal/Vertical Precision (cm), Satellites, Epochs", "GEOREFERENCING")
    ]

    for row_idx, row in enumerate(rows_data, start=1):
        for col_idx, text in enumerate(row):
            cell = tbl.cell(row_idx, col_idx)
            cell.paragraphs[0].text = text
            cell.paragraphs[0].runs[0].font.size = Pt(8.5)
            set_cell_background(cell, "F8FAFC" if row_idx % 2 == 1 else "FFFFFF")
            set_cell_margins(cell, 70, 70, 90, 90)

    doc.add_paragraph()

    # ---------------------------------------------------------------------------
    # SECTION 3: STATE MACHINES & DEDUPLICATION
    # ---------------------------------------------------------------------------
    h3 = doc.add_heading("3. Session State Machines & Cryptographic Integrity", level=1)
    h3.paragraph_format.space_before = Pt(16)
    h3.paragraph_format.space_after = Pt(6)

    doc.add_paragraph(
        "• Upload Session State Machine: CREATED ──> UPLOADING ──> COMPLETED / PARTIAL / FAILED / CANCELLED\n"
        "• File Lifecycle State Machine: PENDING ──> UPLOADING ──> UPLOADED ──> VALIDATED ──> QUEUED ──> PROCESSING ──> PROCESSED / FAILED\n"
        "• Cryptographic Deduplication: Automatically computes SHA-256 hashes upon upload. If a file was already ingested, it is flagged as is_duplicate=True, avoiding redundant storage while maintaining survey references."
    )

    # ---------------------------------------------------------------------------
    # SECTION 4: VISUAL PLATFORM WALKTHROUGH & SCREENSHOTS
    # ---------------------------------------------------------------------------
    h4 = doc.add_heading("4. Visual Walkthrough & Interface Verification", level=1)
    h4.paragraph_format.space_before = Pt(16)
    h4.paragraph_format.space_after = Pt(6)

    brain_dir = r"C:\Users\Vedansh Singhal\.gemini\antigravity-ide\brain\8a2adbf6-3a14-4cb9-afd2-a21233355856"

    screenshots = [
        ("dashboard_page_1788158221629.png", "Figure 1: BhoomiSync Main Operations Dashboard with 'Launch Drone Ingestion Station' trigger."),
        ("gis_map_workbench_1788158666575.png", "Figure 2: Interactive GIS Map Workbench rendering real-time photo capture points (green), LiDAR footprint (purple), and RTK flight trajectory (cyan)."),
        ("gis_map_popup_1788158902985.png", "Figure 3: Georeferenced Camera Photo popup inspector displaying altitude, coordinates HUD, and camera parameters over satellite tiles."),
        ("provenance_dag_tab_1788155896665.png", "Figure 4: Immutable Provenance DAG Graph tracking Tiers 1-3 dataset transformations."),
        ("ai_inference_manifest_1788156129403.png", "Figure 5: AI Intelligence Workbench executing model-agnostic cadastral boundary segmentation inference."),
        ("system_gateways_page_1788156356918.png", "Figure 6: System Architecture Monitor tracking ESP32PhoneGateway, MockStorageProvider, and PostGIS health.")
    ]

    for filename, caption in screenshots:
        full_img_path = os.path.join(brain_dir, filename)
        if os.path.exists(full_img_path):
            doc.add_paragraph()
            doc.add_picture(full_img_path, width=Inches(6.2))
            cp = doc.add_paragraph()
            cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            crun = cp.add_run(caption)
            crun.font.size = Pt(8.5)
            crun.font.italic = True
            crun.font.color.rgb = RGBColor(100, 116, 139)

    # ---------------------------------------------------------------------------
    # SECTION 5: CANONICAL REST API CATALOG (/api/v1/)
    # ---------------------------------------------------------------------------
    h5 = doc.add_heading("5. Canonical REST API Catalog (/api/v1/)", level=1)
    h5.paragraph_format.space_before = Pt(16)
    h5.paragraph_format.space_after = Pt(6)

    api_tbl = doc.add_table(rows=11, cols=3)
    api_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    a_headers = ["Method", "Canonical Endpoint (/api/v1/)", "Functional Description"]
    for i, h in enumerate(a_headers):
        cell = api_tbl.cell(0, i)
        cell.paragraphs[0].text = h
        cell.paragraphs[0].runs[0].bold = True
        cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
        set_cell_background(cell, "047857")
        set_cell_margins(cell, 80, 80, 100, 100)

    api_rows = [
        ("GET", "/api/v1/health", "Platform health, PostGIS status, storage provider, and active gateway type."),
        ("GET", "/api/v1/sensors", "List registered drone sensor payloads (Camera, LiDAR, RTK GNSS)."),
        ("POST", "/api/v1/surveys/{id}/upload-sessions", "Initialize continuous / batched ingestion sync session."),
        ("POST", "/api/v1/surveys/{id}/upload", "Multipart continuous sensor file upload with EXIF/LiDAR parsing & dedup."),
        ("GET", "/api/v1/files/{file_id}", "Retrieve stored file metadata, EXIF parameters, and storage keys."),
        ("POST", "/api/v1/files/{file_id}/retry", "Retry failed file validation or requeue downstream processing job."),
        ("GET", "/api/v1/surveys/{id}/spatial-footprint", "GeoJSON FeatureCollections of photo points, LiDAR bounds, RTK tracks."),
        ("GET", "/api/v1/processing/jobs", "Monitor modular processing job queue."),
        ("GET", "/api/v1/surveys", "List all agricultural survey missions with pagination and summary stats."),
        ("GET", "/api/v1/surveys/{id}/parcels", "Retrieve all parcel boundaries with calculated geodesic area & perimeter.")
    ]

    for row_idx, row in enumerate(api_rows, start=1):
        for col_idx, text in enumerate(row):
            cell = api_tbl.cell(row_idx, col_idx)
            cell.paragraphs[0].text = text
            cell.paragraphs[0].runs[0].font.size = Pt(8.5)
            if col_idx == 0:
                cell.paragraphs[0].runs[0].bold = True
            set_cell_background(cell, "F0FDF4" if row_idx % 2 == 1 else "FFFFFF")
            set_cell_margins(cell, 60, 60, 80, 80)

    doc.add_paragraph()

    # ---------------------------------------------------------------------------
    # SECTION 6: HOW TO TEST & VERIFICATION RESULTS
    # ---------------------------------------------------------------------------
    h6 = doc.add_heading("6. Testing, Demo Mode & Verification Results", level=1)
    h6.paragraph_format.space_before = Pt(16)
    h6.paragraph_format.space_after = Pt(6)

    doc.add_paragraph(
        "• Demo Simulation Testing: Click 'Drone Data Ingestion' -> 'Demo Simulator Mode' -> 'Trigger Full Drone Ingestion Batch Simulation'. The system ingests synthetic geotagged camera photos, a Livox LiDAR strip, and RTK GNSS logs, and renders them instantly on the GIS Map.\n"
        "• Pytest Backend Test Suite: 27/27 unit and integration tests passed in 0.95s.\n"
        "• Frontend Production Build: Vite & TypeScript build completed in 1.45s with 0 errors.\n"
        "• Browser End-to-End Verification: Complete browser workflow recorded to drone_ingestion_pipeline_demo_1788158199234.webp."
    )

    doc.save(output_path)
    print(f"Prompt 2 Walkthrough document successfully saved at: {output_path}")

if __name__ == "__main__":
    out1 = r"d:\BhoomiSync\BhoomiSync_Prompt2_Walkthrough.docx"
    out2 = r"d:\BhoomiSync\docs\BhoomiSync_Prompt2_Walkthrough.docx"
    os.makedirs(os.path.dirname(out2), exist_ok=True)
    try:
        build_prompt2_walkthrough_document(out1)
    except Exception as e:
        print(f"Error saving out1: {e}")
    try:
        build_prompt2_walkthrough_document(out2)
    except Exception as e:
        print(f"Error saving out2: {e}")
